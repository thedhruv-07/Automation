"""One-off importer: converts a BIS FMCS (Foreign Manufacturers Certification
Scheme) register export into clients_certifications.xlsx's single-sheet
client roster schema.

Source format (Manak Online's FMCS "CML No" report, same shape used by
renewals_import.py for the validity-recheck feature): one sheet, columns
found by header name. CML No is used for both Client ID and Certification
ID, since it's the natural unique key in this dataset (mirrors
import_bis_isi_data.py's use of its own licence number). A CML No is listed
once per Standard No it covers, so the same manufacturer routinely has
several rows -- these get separate roster rows (one per certification), with
the Client ID suffixed to stay unique when a CML No repeats. There's no
issue/grant date or phone column in this source, so those are left blank.
Status is recomputed from Validity Date using this project's own urgency
thresholds -- the source "Status" column (Operative/Cancelled/etc.) means
something different and is not used.

Usage: python import_fmcs.py "<path to source xlsx>"
"""
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

from import_helpers import RowCollector, parse_validity_date, compute_status, header_index_map, get

OUTPUT_HEADERS = [
    "Client ID", "Full Name", "Company", "Email", "Phone (WhatsApp)",
    "Certification Name", "Scheme", "Certification ID", "Issue Date", "Expiry Date",
    "Renewal Link", "Status", "Product Name",
]


def looks_like_fmcs_workbook(wb) -> bool:
    """True if any sheet's header row has both a CML-no and name column."""
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            header_row = next(ws.iter_rows(values_only=True))
        except StopIteration:
            continue
        index_map = header_index_map(header_row)
        if "cml no" in index_map and "name" in index_map:
            return True
    return False


def import_fmcs_workbook(wb, out_ws, today=None):
    """Appends converted roster rows from an already-open FMCS workbook into out_ws."""
    today = today or datetime.today().replace(hour=0, minute=0, second=0, microsecond=0)

    seen_client_ids = set()
    stats = {"sheets_used": 0, "sheets_empty": 0, "rows_written": 0, "rows_skipped_missing_key": 0}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            stats["sheets_empty"] += 1
            continue

        index_map = header_index_map(header_row)
        sheet_had_rows = False

        for row in rows_iter:
            if not row or row[0] is None:
                continue
            cml_no = get(row, index_map, "cml no")
            name = get(row, index_map, "name")
            email = get(row, index_map, "email id", "email")
            validity_raw = get(row, index_map, "validity date", "validity")
            cert_name = get(row, index_map, "standard no", "standard")
            product = get(row, index_map, "product")

            if not cml_no or not validity_raw or not cert_name:
                stats["rows_skipped_missing_key"] += 1
                continue

            expiry_dt = parse_validity_date(validity_raw)
            if expiry_dt is None:
                stats["rows_skipped_missing_key"] += 1
                continue

            cert_name = str(cert_name).strip()
            client_id = str(cml_no).strip()
            if client_id in seen_client_ids:
                client_id = f"{client_id}-{cert_name}"
            seen_client_ids.add(client_id)

            out_ws.append([
                client_id,
                name,
                name,
                email,
                None,
                cert_name,
                "FMCS",
                str(cml_no).strip(),
                None,
                expiry_dt.strftime("%d-%m-%Y"),
                None,
                compute_status(expiry_dt, today),
                product,
            ])
            stats["rows_written"] += 1
            sheet_had_rows = True

        if sheet_had_rows:
            stats["sheets_used"] += 1
        else:
            stats["sheets_empty"] += 1

    return stats


def import_fmcs(source_path: str, output_path: str):
    wb = openpyxl.load_workbook(source_path, read_only=True, data_only=True)

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.append(OUTPUT_HEADERS)

    stats = import_fmcs_workbook(wb, out_ws)

    wb.close()
    out_wb.save(output_path)
    return stats


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python import_fmcs.py \"<path to source xlsx>\"")
        sys.exit(1)

    source = sys.argv[1]
    output = str(Path(__file__).parent.parent.parent / "data" / "clients_certifications.xlsx")
    result = import_fmcs(source, output)
    print("Sheets with data used:", result["sheets_used"])
    print("Sheets empty/skipped:", result["sheets_empty"])
    print("Rows written:", result["rows_written"])
    print("Rows skipped (missing CML no / validity date / standard no):", result["rows_skipped_missing_key"])
    print("Saved to:", output)
